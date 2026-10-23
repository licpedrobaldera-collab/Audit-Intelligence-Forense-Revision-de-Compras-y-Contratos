from __future__ import annotations
import os, re, hashlib, hmac, sqlite3, logging, mimetypes
from pathlib import Path
from datetime import datetime, date
import streamlit as st
import pandas as pd

BASE_DIR = Path(__file__).resolve().parent
DB_FILE = BASE_DIR / "database" / "audit_intelligence.db"
UPLOAD_DIR, EXPORT_DIR, LOG_DIR = BASE_DIR / "uploads", BASE_DIR / "exports", BASE_DIR / "logs"
for folder in (DB_FILE.parent, UPLOAD_DIR, EXPORT_DIR, LOG_DIR):
    folder.mkdir(parents=True, exist_ok=True)
logging.basicConfig(filename=LOG_DIR / "app.log", level=logging.INFO,
                    format="%(asctime)s %(levelname)s %(message)s")
APP_NAME = os.getenv("AI_APP_NAME", "Audit Intelligence USA")

def connect():
    conn = sqlite3.connect(DB_FILE, timeout=20)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys=ON")
    conn.execute("PRAGMA journal_mode=WAL")
    return conn

def init_db():
    schema = """
    CREATE TABLE IF NOT EXISTS users(id INTEGER PRIMARY KEY AUTOINCREMENT, username TEXT UNIQUE NOT NULL, password_hash TEXT NOT NULL, role TEXT NOT NULL DEFAULT 'Auditor', active INTEGER NOT NULL DEFAULT 1, created_at TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS expedientes(id INTEGER PRIMARY KEY AUTOINCREMENT, codigo TEXT UNIQUE NOT NULL, nombre TEXT NOT NULL, entidad TEXT DEFAULT '', objetivo TEXT DEFAULT '', responsable TEXT DEFAULT '', fecha TEXT NOT NULL, estado TEXT NOT NULL DEFAULT 'Activo', created_by TEXT NOT NULL, created_at TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS hallazgos(id INTEGER PRIMARY KEY AUTOINCREMENT, expediente_id INTEGER, titulo TEXT NOT NULL, condicion TEXT DEFAULT '', criterio TEXT DEFAULT '', causa TEXT DEFAULT '', efecto TEXT DEFAULT '', recomendacion TEXT DEFAULT '', nivel TEXT NOT NULL DEFAULT 'Medio', estado TEXT NOT NULL DEFAULT 'Abierto', created_by TEXT NOT NULL, created_at TEXT NOT NULL, FOREIGN KEY(expediente_id) REFERENCES expedientes(id) ON DELETE SET NULL);
    CREATE TABLE IF NOT EXISTS riesgos(id INTEGER PRIMARY KEY AUTOINCREMENT, expediente_id INTEGER, categoria TEXT NOT NULL, descripcion TEXT NOT NULL, probabilidad INTEGER NOT NULL, impacto INTEGER NOT NULL, puntuacion INTEGER NOT NULL, nivel TEXT NOT NULL, mitigacion TEXT DEFAULT '', estado TEXT DEFAULT 'Pendiente', created_at TEXT NOT NULL, FOREIGN KEY(expediente_id) REFERENCES expedientes(id) ON DELETE SET NULL);
    CREATE TABLE IF NOT EXISTS proveedores(id INTEGER PRIMARY KEY AUTOINCREMENT, nombre TEXT NOT NULL, rnc TEXT DEFAULT '', contacto TEXT DEFAULT '', estado TEXT DEFAULT 'Activo', created_at TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS contratos(id INTEGER PRIMARY KEY AUTOINCREMENT, numero TEXT UNIQUE NOT NULL, proveedor_id INTEGER, objeto TEXT DEFAULT '', monto REAL NOT NULL DEFAULT 0, fecha_inicio TEXT DEFAULT '', fecha_fin TEXT DEFAULT '', estado TEXT DEFAULT 'Vigente', created_at TEXT NOT NULL, FOREIGN KEY(proveedor_id) REFERENCES proveedores(id) ON DELETE SET NULL);
    CREATE TABLE IF NOT EXISTS pagos(id INTEGER PRIMARY KEY AUTOINCREMENT, contrato_id INTEGER, beneficiario TEXT NOT NULL, referencia TEXT DEFAULT '', monto REAL NOT NULL CHECK(monto >= 0), fecha TEXT NOT NULL, descripcion TEXT DEFAULT '', created_at TEXT NOT NULL, FOREIGN KEY(contrato_id) REFERENCES contratos(id) ON DELETE SET NULL);
    CREATE TABLE IF NOT EXISTS evidencias(id INTEGER PRIMARY KEY AUTOINCREMENT, expediente_id INTEGER, nombre_original TEXT NOT NULL, nombre_guardado TEXT NOT NULL, ruta TEXT NOT NULL, sha256 TEXT NOT NULL, tipo TEXT DEFAULT '', texto_extraido TEXT DEFAULT '', subido_por TEXT NOT NULL, created_at TEXT NOT NULL, FOREIGN KEY(expediente_id) REFERENCES expedientes(id) ON DELETE SET NULL);
    CREATE TABLE IF NOT EXISTS actividad(id INTEGER PRIMARY KEY AUTOINCREMENT, usuario TEXT NOT NULL, accion TEXT NOT NULL, detalle TEXT DEFAULT '', fecha TEXT NOT NULL);
    """
    with connect() as conn:
        conn.executescript(schema)

def log_action(user, action, detail=""):
    try:
        with connect() as conn:
            conn.execute("INSERT INTO actividad(usuario,accion,detalle,fecha) VALUES(?,?,?,?)",
                (user or "sistema", action, str(detail)[:1000], datetime.now().isoformat(timespec="seconds")))
    except Exception:
        logging.exception("No se pudo registrar actividad")

def count(table):
    allowed = {"users","expedientes","hallazgos","riesgos","proveedores","contratos","pagos","evidencias"}
    if table not in allowed: return 0
    with connect() as conn: return conn.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]

def fetch_df(sql, params=()):
    with connect() as conn: return pd.read_sql_query(sql, conn, params=params)

def password_hash(password):
    import secrets
    salt = secrets.token_bytes(16)
    digest = hashlib.pbkdf2_hmac("sha256", password.encode(), salt, 310000)
    return f"pbkdf2_sha256$310000${salt.hex()}${digest.hex()}"

def verify_password(password, encoded):
    try:
        alg, rounds, salt, expected = encoded.split("$")
        actual = hashlib.pbkdf2_hmac("sha256", password.encode(), bytes.fromhex(salt), int(rounds)).hex()
        return alg == "pbkdf2_sha256" and hmac.compare_digest(actual, expected)
    except Exception: return False

def ensure_initial_admin():
    with connect() as conn:
        if conn.execute("SELECT COUNT(*) FROM users").fetchone()[0] == 0:
            password = os.getenv("AI_INITIAL_ADMIN_PASSWORD", "Cambiar-Admin-2026!")
            conn.execute("INSERT INTO users(username,password_hash,role,active,created_at) VALUES(?,?,?,?,?)",
                ("admin", password_hash(password), "Administrador", 1, datetime.now().isoformat(timespec="seconds")))
            (BASE_DIR / "CAMBIAR_CREDENCIALES.txt").write_text(
                f"Usuario inicial: admin\nContraseña inicial: {password}\nCambie la clave inmediatamente desde Seguridad y elimine este archivo.\n", encoding="utf-8")

def extract_text(path):
    ext = path.suffix.lower()
    try:
        if ext == ".txt": return path.read_text(encoding="utf-8", errors="replace")
        if ext == ".csv": return pd.read_csv(path, dtype=str).fillna("").to_string(index=False)
        if ext in (".xlsx", ".xls"):
            xls = pd.ExcelFile(path)
            return "\n\n".join(f"HOJA: {s}\n{pd.read_excel(path, sheet_name=s, dtype=str).fillna('').to_string(index=False)}" for s in xls.sheet_names)
        if ext == ".docx":
            from docx import Document
            doc = Document(path)
            return "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(
                " | ".join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows)
        if ext == ".pdf":
            try:
                import pymupdf
                with pymupdf.open(path) as pdf: return "\n".join(page.get_text() for page in pdf)
            except ImportError:
                import pypdf
                return "\n".join(page.extract_text() or "" for page in pypdf.PdfReader(str(path)).pages)
    except Exception as exc:
        logging.exception("Error extrayendo %s", path.name)
        return f"[No se pudo extraer texto: {type(exc).__name__}: {exc}]"
    return "[Archivo guardado. Este formato no tiene extractor de texto habilitado.]"

def ai_analyze(text):
    """Motor local explicable basado en reglas; no simula ser un LLM."""
    text = (text or "").lower()
    rules = [
        ("Posible pago duplicado", ["duplicado","misma factura","doble pago","repetido"]),
        ("Riesgo de contratación", ["contrato","adenda","licitación","proveedor","compra"]),
        ("Falta de evidencia documental", ["sin evidencia","no consta","no se encontró","incompleto","ausente"]),
        ("Riesgo de plazo", ["vencido","fuera de plazo","fecha límite","retraso"]),
        ("Riesgo de conciliación financiera", ["diferencia","no concilia","monto distinto","descuadre"]),
        ("Riesgo de cumplimiento", ["incumplimiento","normativa","procedimiento","autorización"]),
    ]
    found = [{"riesgo": label, "indicios": [word for word in words if word in text]}
             for label, words in rules if any(word in text for word in words)]
    if not found: found = [{"riesgo":"Revisión manual requerida","indicios":["No se detectaron términos de las reglas básicas"]}]
    return {"motor":"Reglas locales explicables; no es IA generativa",
            "riesgos_detectados":found,
            "limitacion":"Indicadores preliminares, no conclusiones. Validar con documentos originales, criterios y normativa vigente."}

def download_df(df, filename):
    st.download_button("Descargar CSV", df.to_csv(index=False).encode("utf-8-sig"),
                       file_name=filename + ".csv", mime="text/csv")

st.set_page_config(page_title=APP_NAME, page_icon="🔎", layout="wide")
init_db()
ensure_initial_admin()
if "auth_user" not in st.session_state: st.session_state.auth_user = None
if "auth_role" not in st.session_state: st.session_state.auth_role = None

def login_screen():
    st.title("🔎 " + APP_NAME)
    st.caption("Control de auditoría, expedientes, riesgos y evidencias")
    with st.form("login_form"):
        username = st.text_input("Usuario")
        password = st.text_input("Contraseña", type="password")
        submit = st.form_submit_button("Ingresar", type="primary")
    if submit:
        with connect() as conn:
            row = conn.execute("SELECT username,password_hash,role,active FROM users WHERE username=?",
                               (username.strip(),)).fetchone()
        if row and row["active"] and verify_password(password, row["password_hash"]):
            st.session_state.auth_user, st.session_state.auth_role = row["username"], row["role"]
            log_action(username.strip(), "LOGIN", "Acceso correcto")
            st.rerun()
        else: st.error("Credenciales incorrectas o usuario inactivo.")
    st.info("Instalación nueva: usuario admin y clave inicial indicada en CAMBIAR_CREDENCIALES.txt. Cámbiela inmediatamente.")

if not st.session_state.auth_user:
    login_screen()
    st.stop()

user, role = st.session_state.auth_user, st.session_state.auth_role
st.sidebar.title("Audit Intelligence USA")
st.sidebar.caption(f"{user} · {role}")
menu = st.sidebar.radio("Módulos", ["Dashboard","Expedientes","Hallazgos","Riesgos","Proveedores","Contratos","Pagos","Documentos e IA","Reportes","Seguridad"])
if st.sidebar.button("Cerrar sesión"):
    log_action(user,"LOGOUT")
    st.session_state.auth_user = st.session_state.auth_role = None
    st.rerun()

def expediente_options():
    df = fetch_df("SELECT id,codigo,nombre FROM expedientes ORDER BY id DESC")
    options = [None] + df.id.tolist()
    def label(value):
        if value is None: return "Sin expediente"
        row = df[df.id == value].iloc[0]
        return f"{row['codigo']} — {row['nombre']}"
    return options, label

if menu == "Dashboard":
    st.title("Dashboard ejecutivo")
    cols = st.columns(5)
    for col, label, table in zip(cols,["Expedientes","Hallazgos","Riesgos","Proveedores","Evidencias"],
                                 ["expedientes","hallazgos","riesgos","proveedores","evidencias"]):
        col.metric(label,count(table))
    st.subheader("Hallazgos recientes")
    st.dataframe(fetch_df("SELECT id,titulo,nivel,estado,created_at FROM hallazgos ORDER BY id DESC LIMIT 10"),
                 use_container_width=True,hide_index=True)
    risks=fetch_df("SELECT nivel,COUNT(*) cantidad FROM riesgos GROUP BY nivel")
    st.subheader("Riesgos por nivel")
    if not risks.empty: st.bar_chart(risks.set_index("nivel"))
    else: st.caption("No hay riesgos registrados todavía.")

elif menu == "Expedientes":
    st.title("Gestión de expedientes")
    with st.form("expediente_form"):
        codigo=st.text_input("Código único",value=f"AUD-{date.today():%Y%m%d}")
        nombre=st.text_input("Nombre del expediente *")
        entidad=st.text_input("Entidad / organización")
        objetivo=st.text_area("Objetivo y alcance")
        responsable=st.text_input("Responsable")
        estado=st.selectbox("Estado",["Activo","En revisión","Cerrado","Suspendido"])
        submit=st.form_submit_button("Guardar expediente")
    if submit:
        if not codigo.strip() or not nombre.strip(): st.error("Código y nombre son obligatorios.")
        else:
            try:
                with connect() as conn:
                    conn.execute("INSERT INTO expedientes(codigo,nombre,entidad,objetivo,responsable,fecha,estado,created_by,created_at) VALUES(?,?,?,?,?,?,?,?,?)",
                        (codigo.strip(),nombre.strip(),entidad,objetivo,responsable,date.today().isoformat(),estado,user,datetime.now().isoformat(timespec="seconds")))
                log_action(user,"CREAR_EXPEDIENTE",codigo); st.success("Expediente guardado.")
            except sqlite3.IntegrityError: st.error("El código del expediente ya existe.")
    st.dataframe(fetch_df("SELECT * FROM expedientes ORDER BY id DESC"),use_container_width=True,hide_index=True)

elif menu == "Hallazgos":
    st.title("Hallazgos de auditoría")
    options, label = expediente_options()
    with st.form("finding_form"):
        exp=st.selectbox("Expediente",options,format_func=label)
        titulo=st.text_input("Título *"); condicion=st.text_area("Condición / situación observada")
        criterio=st.text_area("Criterio / norma o procedimiento"); causa=st.text_area("Causa")
        efecto=st.text_area("Efecto / riesgo"); recomendacion=st.text_area("Recomendación")
        nivel=st.selectbox("Nivel",["Bajo","Medio","Alto","Crítico"])
        estado=st.selectbox("Estado",["Abierto","En proceso","Cerrado"])
        submit=st.form_submit_button("Registrar hallazgo")
    if submit:
        if not titulo.strip(): st.error("El título es obligatorio.")
        else:
            with connect() as conn:
                conn.execute("INSERT INTO hallazgos(expediente_id,titulo,condicion,criterio,causa,efecto,recomendacion,nivel,estado,created_by,created_at) VALUES(?,?,?,?,?,?,?,?,?,?,?)",
                    (exp,titulo,condicion,criterio,causa,efecto,recomendacion,nivel,estado,user,datetime.now().isoformat(timespec="seconds")))
            log_action(user,"CREAR_HALLAZGO",titulo); st.success("Hallazgo registrado.")
    df=fetch_df("SELECT h.id,e.codigo expediente,h.titulo,h.nivel,h.estado,h.created_at FROM hallazgos h LEFT JOIN expedientes e ON e.id=h.expediente_id ORDER BY h.id DESC")
    st.dataframe(df,use_container_width=True,hide_index=True); download_df(df,"hallazgos")

elif menu == "Riesgos":
    st.title("Matriz de riesgos")
    options,label=expediente_options()
    with st.form("risk_form"):
        exp=st.selectbox("Expediente",options,format_func=label)
        categoria=st.selectbox("Categoría",["Financiero","Cumplimiento","Operacional","Tecnológico","Contratación","Documental","Fraude potencial","Otro"])
        descripcion=st.text_area("Descripción del riesgo *")
        prob=st.slider("Probabilidad",1,5,3); impacto=st.slider("Impacto",1,5,3)
        mitigacion=st.text_area("Control / mitigación propuesta")
        submit=st.form_submit_button("Evaluar y guardar")
    if submit:
        if not descripcion.strip(): st.error("La descripción es obligatoria.")
        else:
            score=prob*impact
            nivel="Crítico" if score>=20 else "Alto" if score>=15 else "Medio" if score>=8 else "Bajo"
            with connect() as conn:
                conn.execute("INSERT INTO riesgos(expediente_id,categoria,descripcion,probabilidad,impacto,puntuacion,nivel,mitigacion,created_at) VALUES(?,?,?,?,?,?,?,?,?)",
                    (exp,categoria,descripcion,prob,impacto,score,nivel,mitigacion,datetime.now().isoformat(timespec="seconds")))
            log_action(user,"CREAR_RIESGO",f"{categoria}: {score}"); st.success(f"Riesgo guardado: {score}/25 — {nivel}")
    df=fetch_df("SELECT * FROM riesgos ORDER BY puntuacion DESC")
    st.dataframe(df,use_container_width=True,hide_index=True); download_df(df,"matriz_riesgos")

elif menu == "Proveedores":
    st.title("Proveedores")
    with st.form("supplier_form"):
        nombre=st.text_input("Nombre / razón social *"); rnc=st.text_input("RNC / identificación")
        contacto=st.text_input("Contacto"); estado=st.selectbox("Estado",["Activo","Inactivo","En revisión"])
        submit=st.form_submit_button("Guardar proveedor")
    if submit:
        if not nombre.strip(): st.error("El nombre es obligatorio.")
        else:
            with connect() as conn: conn.execute("INSERT INTO proveedores(nombre,rnc,contacto,estado,created_at) VALUES(?,?,?,?,?)",(nombre,rnc,contacto,estado,datetime.now().isoformat(timespec="seconds")))
            log_action(user,"CREAR_PROVEEDOR",nombre); st.success("Proveedor registrado.")
    df=fetch_df("SELECT * FROM proveedores ORDER BY id DESC"); st.dataframe(df,use_container_width=True,hide_index=True); download_df(df,"proveedores")

elif menu == "Contratos":
    st.title("Contratos")
    prov=fetch_df("SELECT id,nombre FROM proveedores ORDER BY nombre")
    with st.form("contract_form"):
        numero=st.text_input("Número de contrato *")
        proveedor=st.selectbox("Proveedor",[None]+prov.id.tolist(),format_func=lambda x:"Sin proveedor" if x is None else prov.loc[prov.id==x,"nombre"].iloc[0])
        objeto=st.text_area("Objeto"); monto=st.number_input("Monto",min_value=0.0,step=1000.0)
        inicio=st.date_input("Fecha de inicio",value=date.today()); fin=st.date_input("Fecha de fin",value=date.today())
        estado=st.selectbox("Estado",["Vigente","Vencido","Cerrado","Suspendido"])
        submit=st.form_submit_button("Guardar contrato")
    if submit:
        if not numero.strip(): st.error("El número es obligatorio.")
        elif fin<inicio: st.error("La fecha final no puede ser anterior a la inicial.")
        else:
            try:
                with connect() as conn: conn.execute("INSERT INTO contratos(numero,proveedor_id,objeto,monto,fecha_inicio,fecha_fin,estado,created_at) VALUES(?,?,?,?,?,?,?,?)",(numero,proveedor,objeto,monto,inicio.isoformat(),fin.isoformat(),estado,datetime.now().isoformat(timespec="seconds")))
                log_action(user,"CREAR_CONTRATO",numero); st.success("Contrato guardado.")
            except sqlite3.IntegrityError: st.error("Ese número de contrato ya existe.")
    st.dataframe(fetch_df("SELECT c.*,p.nombre proveedor FROM contratos c LEFT JOIN proveedores p ON p.id=c.proveedor_id ORDER BY c.id DESC"),use_container_width=True,hide_index=True)

elif menu == "Pagos":
    st.title("Registro y revisión de pagos")
    contracts=fetch_df("SELECT id,numero FROM contratos ORDER BY numero")
    with st.form("payment_form"):
        contrato=st.selectbox("Contrato",[None]+contracts.id.tolist(),format_func=lambda x:"Sin contrato" if x is None else contracts.loc[contracts.id==x,"numero"].iloc[0])
        beneficiario=st.text_input("Beneficiario *"); referencia=st.text_input("Referencia / factura")
        monto=st.number_input("Monto",min_value=0.0,step=100.0); fecha=st.date_input("Fecha",value=date.today())
        descripcion=st.text_area("Descripción"); submit=st.form_submit_button("Registrar pago")
    if submit:
        if not beneficiario.strip() or monto<=0: st.error("Indique beneficiario y monto mayor que cero.")
        else:
            with connect() as conn: conn.execute("INSERT INTO pagos(contrato_id,beneficiario,referencia,monto,fecha,descripcion,created_at) VALUES(?,?,?,?,?,?,?)",(contrato,beneficiario,referencia,monto,fecha.isoformat(),descripcion,datetime.now().isoformat(timespec="seconds")))
            log_action(user,"CREAR_PAGO",f"{beneficiario} {monto}"); st.success("Pago registrado.")
    df=fetch_df("SELECT p.*,c.numero contrato FROM pagos p LEFT JOIN contratos c ON c.id=p.contrato_id ORDER BY p.id DESC")
    st.dataframe(df,use_container_width=True,hide_index=True)
    if not df.empty:
        dup=df[df.duplicated(subset=["referencia","monto","beneficiario"],keep=False) & (df.referencia.fillna("")!="")]
        if not dup.empty: st.warning("Posibles pagos duplicados: valide contra facturas y comprobantes."); st.dataframe(dup,use_container_width=True,hide_index=True)
    download_df(df,"pagos")

elif menu == "Documentos e IA":
    st.title("Evidencias documentales e IA asistida")
    options,label=expediente_options()
    exp=st.selectbox("Vincular al expediente",options,format_func=label)
    files=st.file_uploader("Cargar evidencias (máx. 25 MB cada una)",type=["pdf","docx","xlsx","xls","csv","txt","png","jpg","jpeg"],accept_multiple_files=True)
    if files:
        for file in files:
            raw=file.getvalue()
            if len(raw)>25*1024*1024: st.error(f"{file.name}: supera 25 MB."); continue
            digest=hashlib.sha256(raw).hexdigest()
            safe=re.sub(r"[^A-Za-z0-9._-]","_",Path(file.name).name)[:120]
            saved=f"{datetime.now():%Y%m%d%H%M%S}_{digest[:10]}_{safe}"
            destination=UPLOAD_DIR/saved
            if not destination.exists(): destination.write_bytes(raw)
            text=extract_text(destination)
            with connect() as conn:
                exists=conn.execute("SELECT id FROM evidencias WHERE sha256=?",(digest,)).fetchone()
                if not exists:
                    conn.execute("INSERT INTO evidencias(expediente_id,nombre_original,nombre_guardado,ruta,sha256,tipo,texto_extraido,subido_por,created_at) VALUES(?,?,?,?,?,?,?,?,?)",
                        (exp,file.name,saved,str(destination.relative_to(BASE_DIR)),digest,file.type or mimetypes.guess_type(file.name)[0] or "",text,user,datetime.now().isoformat(timespec="seconds")))
                    log_action(user,"CARGAR_EVIDENCIA",file.name)
            with st.expander(f"Documento: {file.name}"):
                st.caption(f"SHA-256: {digest}")
                st.text_area("Texto extraído — verificar con el original",text,height=180,key=f"text_{digest}")
                if st.button("Analizar con IA local",key=f"analyze_{digest}"):
                    st.json(ai_analyze(text)); log_action(user,"ANALISIS_IA",file.name)
    st.subheader("Evidencias registradas")
    df=fetch_df("SELECT id,expediente_id,nombre_original,sha256,subido_por,created_at FROM evidencias ORDER BY id DESC")
    st.dataframe(df,use_container_width=True,hide_index=True)
    st.caption("PDF escaneados e imágenes requieren OCR adicional. SHA-256 ayuda a verificar integridad del archivo, pero no demuestra autenticidad.")

elif menu == "Reportes":
    st.title("Reportes y exportaciones")
    datasets={"Expedientes":"SELECT * FROM expedientes","Hallazgos":"SELECT * FROM hallazgos","Riesgos":"SELECT * FROM riesgos",
      "Proveedores":"SELECT * FROM proveedores","Contratos":"SELECT * FROM contratos","Pagos":"SELECT * FROM pagos",
      "Evidencias":"SELECT id,expediente_id,nombre_original,sha256,subido_por,created_at FROM evidencias","Actividad":"SELECT * FROM actividad"}
    selected=st.selectbox("Datos a exportar",list(datasets))
    df=fetch_df(datasets[selected]); st.dataframe(df,use_container_width=True,hide_index=True)
    download_df(df,selected.lower().replace(" ","_"))
    if st.button("Crear copia de seguridad"):
        import shutil
        folder=BASE_DIR/"backups"; folder.mkdir(exist_ok=True)
        backup=folder/f"audit_backup_{datetime.now():%Y%m%d_%H%M%S}.db"
        with connect() as source, sqlite3.connect(backup) as target: source.backup(target)
        st.success(f"Copia creada: {backup.name}"); log_action(user,"BACKUP",backup.name)

elif menu == "Seguridad":
    st.title("Seguridad y usuarios")
    with st.form("password_form"):
        old=st.text_input("Contraseña actual",type="password")
        new=st.text_input("Nueva contraseña (mínimo 12 caracteres)",type="password")
        confirm=st.text_input("Confirmar contraseña",type="password")
        submit=st.form_submit_button("Cambiar contraseña")
    if submit:
        if len(new)<12: st.error("Use al menos 12 caracteres.")
        elif new!=confirm: st.error("Las contraseñas no coinciden.")
        else:
            with connect() as conn: row=conn.execute("SELECT password_hash FROM users WHERE username=?",(user,)).fetchone()
            if not row or not verify_password(old,row["password_hash"]): st.error("La contraseña actual no es correcta.")
            else:
                with connect() as conn: conn.execute("UPDATE users SET password_hash=? WHERE username=?",(password_hash(new),user))
                log_action(user,"CAMBIAR_PASSWORD"); st.success("Contraseña actualizada.")
    if role=="Administrador":
        st.subheader("Crear usuario")
        with st.form("create_user_form"):
            username=st.text_input("Usuario nuevo")
            password=st.text_input("Contraseña temporal (mínimo 12 caracteres)",type="password")
            new_role=st.selectbox("Rol",["Auditor","Supervisor","Administrador"])
            submit_user=st.form_submit_button("Crear usuario")
        if submit_user:
            if not re.fullmatch(r"[A-Za-z0-9_.-]{3,40}",username): st.error("Usuario inválido: use 3-40 letras, números, punto, guion o guion bajo.")
            elif len(password)<12: st.error("La contraseña debe tener al menos 12 caracteres.")
            else:
                try:
                    with connect() as conn: conn.execute("INSERT INTO users(username,password_hash,role,active,created_at) VALUES(?,?,?,?,?)",(username,password_hash(password),new_role,1,datetime.now().isoformat(timespec="seconds")))
                    log_action(user,"CREAR_USUARIO",username); st.success("Usuario creado.")
                except sqlite3.IntegrityError: st.error("El usuario ya existe.")
        st.dataframe(fetch_df("SELECT id,username,role,active,created_at FROM users"),use_container_width=True,hide_index=True)
    st.subheader("Registro de actividad")
    st.dataframe(fetch_df("SELECT * FROM actividad ORDER BY id DESC LIMIT 200"),use_container_width=True,hide_index=True)

st.sidebar.caption("Versión base del proyecto · no publicar directamente en Internet sin endurecimiento adicional")
